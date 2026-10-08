import gc
import re
import sys
import tempfile
import tracemalloc
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet._read_only import ReadOnlyWorksheet


PROJECT_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_DIR / "src"))

from spreadsheet_converter import _rows_to_md, write_xlsx_sheets


def save_workbook(path, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def collect_sheets(path, tmp_path):
    written = []

    def write_sheet(body, title, source_file, word_count, output_dir, source_type,
                    vault_dir, header_extras):
        assert not isinstance(body, (str, list, tuple)), "Markdown must be streamed"
        markdown = "".join(body)
        assert word_count == len(markdown.split())
        written.append((markdown, title, source_file, word_count, header_extras))
        return SimpleNamespace(output_path=str(output_dir / f"sheet-{len(written)}.md"))

    result = write_xlsx_sheets(str(path), tmp_path / "out", None, write_sheet)
    return written, result


def test_streamed_markdown_preserves_uneven_rows_blanks_and_escaping(tmp_path):
    path = tmp_path / "Workbook.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name"])
    sheet.append([" a | b\nc ", None, "tail\tword"])
    sheet.append([])
    sheet.append([0, False])
    # Formatting creates actual empty rows and columns in the worksheet XML.
    sheet["F8"].font = Font(bold=True)
    workbook.save(path)
    workbook.close()

    written, result = collect_sheets(path, tmp_path)

    expected = (
        "| Name |  |  |\n"
        "| --- | --- | --- |\n"
        "| a \\| b c |  | tail\tword |\n"
        "|  |  |  |\n"
        "| 0 | False |  |"
    )
    assert written == [(expected, "Workbook Data", "Workbook.xlsx",
                        len(expected.split()), {"Sheet": "Data"})]
    assert result == (1, len(expected.split()), ["sheet-1.md"])


def test_streamed_markdown_keeps_empty_first_row_as_header(tmp_path):
    path = tmp_path / "empty-header.xlsx"
    save_workbook(path, [[], [None, "Value"], [], ["last"]])

    written, _ = collect_sheets(path, tmp_path)

    assert written[0][0] == (
        "|  |  |\n| --- | --- |\n|  | Value |\n|  |  |\n| last |  |"
    )


def test_empty_and_whitespace_only_sheets_are_skipped(tmp_path):
    path = tmp_path / "empty.xlsx"
    workbook = Workbook()
    workbook.create_sheet("Whitespace").append([None, " \n\t "])
    workbook.create_sheet("Formatted")["Z100"].font = Font(bold=True)
    workbook.save(path)
    workbook.close()

    written, result = collect_sheets(path, tmp_path)

    assert written == []
    assert result == (0, 0, [])


def test_multiple_sheets_keep_independent_dimensions_and_word_counts(tmp_path):
    path = tmp_path / "Workbook.xlsx"
    workbook = Workbook()
    workbook.active.title = "Wide"
    workbook.active.append(["one", "two", "three"])
    workbook.create_sheet("Empty")
    workbook.create_sheet("Narrow").append(["single"])
    workbook.save(path)
    workbook.close()

    written, result = collect_sheets(path, tmp_path)

    assert written[0][0] == "| one | two | three |\n| --- | --- | --- |"
    assert written[1][0] == "| single |\n| --- |"
    assert result == (2, sum(len(row[0].split()) for row in written),
                      ["sheet-1.md", "sheet-2.md"])


@pytest.mark.parametrize("dimension", ["A1:XFD1048576", "A1:A1"])
def test_incorrect_dimensions_do_not_pad_or_truncate_actual_rows(
    tmp_path, monkeypatch, dimension
):
    path = tmp_path / "dimensions.xlsx"
    save_workbook(path, [["A", "B", "C"], ["one", "two", "three"]])
    with ZipFile(path) as archive:
        entries = [(item, archive.read(item.filename)) for item in archive.infolist()]
    with ZipFile(path, "w") as archive:
        for item, data in entries:
            if item.filename == "xl/worksheets/sheet1.xml":
                data, count = re.subn(
                    rb'<dimension ref="[^"]+"',
                    f'<dimension ref="{dimension}"'.encode(),
                    data,
                )
                assert count == 1
            archive.writestr(item, data)

    original_iter_rows = ReadOnlyWorksheet.iter_rows

    def bounded_rows(sheet, *args, **kwargs):
        for index, row in enumerate(original_iter_rows(sheet, *args, **kwargs)):
            assert index < 2, "Inflated dimensions generated empty rows"
            assert len(row) <= 3, "Inflated dimensions padded actual rows"
            yield row

    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", bounded_rows)
    written, _ = collect_sheets(path, tmp_path)

    assert written[0][0] == (
        "| A | B | C |\n| --- | --- | --- |\n| one | two | three |"
    )


def test_large_workbook_is_written_without_buffering_full_markdown(tmp_path):
    path = tmp_path / "large.xlsx"
    row_count = 6_000
    payload = "x" * 1_000
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Large")
    sheet.append(["Index", "Payload"])
    for index in range(row_count):
        sheet.append([index, payload])
    workbook.save(path)
    workbook.close()
    del workbook, sheet
    gc.collect()

    def consume_sheet(body, title, source_file, word_count, output_dir, source_type,
                      vault_dir, header_extras):
        char_count = 0
        actual_words = 0
        chunks = (body,) if isinstance(body, str) else body
        for chunk in chunks:
            char_count += len(chunk)
            actual_words += len(chunk.split())
        assert char_count > row_count * len(payload)
        assert actual_words == word_count
        return SimpleNamespace(output_path=str(output_dir / "large.md"))

    tracemalloc.start()
    try:
        result = write_xlsx_sheets(str(path), tmp_path / "out", None, consume_sheet)
        _, peak_bytes = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()

    assert result[0] == 1
    assert peak_bytes < 4_000_000, f"Conversion buffered {peak_bytes:,} bytes"


def test_rows_to_md_compatibility_helper_still_accepts_uneven_rows():
    assert _rows_to_md([["A"], ["B", "C"], [], ["", ""]]) == (
        "| A |  |\n| --- | --- |\n| B | C |"
    )


def test_writer_failure_closes_workbook_and_temporary_spool(tmp_path, monkeypatch):
    import openpyxl

    path = tmp_path / "failure.xlsx"
    save_workbook(path, [["Header"], ["Value"]])
    workbooks = []
    spools = []
    original_load = openpyxl.load_workbook

    def load_workbook(*args, **kwargs):
        workbook = original_load(*args, **kwargs)
        workbooks.append(workbook)
        return workbook

    def temporary_file(*args, **kwargs):
        spool = tempfile.TemporaryFile(*args, **kwargs)
        spools.append(spool)
        return spool

    def fail_after_header(body, *args, **kwargs):
        assert next(iter(body)) == "| Header |"
        raise OSError("Output disk full")

    monkeypatch.setattr(openpyxl, "load_workbook", load_workbook)
    monkeypatch.setattr("spreadsheet_converter.TemporaryFile", temporary_file)
    with pytest.raises(OSError, match="Output disk full"):
        write_xlsx_sheets(str(path), tmp_path / "out", None, fail_after_header)

    assert len(spools) == len(workbooks) == 1
    assert spools[0].closed
    assert workbooks[0]._archive.fp is None
