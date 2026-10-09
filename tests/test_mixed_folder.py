from pathlib import Path
import os
import plistlib
import sys
from unittest import mock

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_cli_mode import load_converter_app


def test_add_folder_stages_mixed_supported_files_recursively(tmp_path):
    app = load_converter_app()
    app.SUPPORTED = {'.pdf', '.docx', '.xlsx', '.txt', '.rtf', '.html', '.htm', '.png', '.jpg', '.jpeg', '.webp'}
    nested = tmp_path / 'nested'
    nested.mkdir()
    names = ['report.DOCX', 'notes.txt', 'book.pdf', 'table.xlsx', 'quote.jpg', 'page.html', 'letter.rtf']
    expected = []
    for index, name in enumerate(names):
        path = (nested if index % 2 else tmp_path) / name
        path.write_bytes(b'fixture')
        expected.append(str(path))
    (tmp_path / 'ignored.zip').write_bytes(b'not supported')
    api = app.Api()
    api.window = mock.Mock()
    api.window.create_file_dialog.return_value = [str(tmp_path)]
    api.add_folder()
    assert set(api._collect_staged_paths()) == set(expected)


def test_dropped_document_folder_is_not_treated_as_empty(tmp_path):
    app = load_converter_app()
    app.SUPPORTED = {'.docx', '.txt'}
    path = tmp_path / 'document.docx'
    path.write_bytes(b'fixture')
    api = app.Api()
    api.stage_files([str(tmp_path), str(path)])
    assert api._collect_staged_paths() == [str(path)]
    assert len(api._staged_folders) == 1


def test_overlapping_mixed_folders_are_deduplicated(tmp_path):
    app = load_converter_app()
    app.SUPPORTED = {'.docx', '.pdf', '.jpg'}
    nested = tmp_path / 'nested'
    nested.mkdir()
    first, second = tmp_path / 'first.docx', nested / 'second.pdf'
    first.write_bytes(b'fixture')
    second.write_bytes(b'fixture')
    api = app.Api()
    api.window = mock.Mock()
    api.window.create_file_dialog.side_effect = [[str(tmp_path)], [str(nested)]]
    api.add_folder()
    api.add_folder()
    assert set(api._collect_staged_paths()) == {str(first), str(second)}
    assert len(api._collect_staged_paths()) == 2


def test_mixed_folder_converts_real_documents_and_spreadsheets(tmp_path, monkeypatch):
    import converter_app as app
    from docx import Document
    from openpyxl import Workbook
    import pymupdf

    source = tmp_path / 'source'
    nested = source / 'nested'
    nested.mkdir(parents=True)
    doc = Document()
    doc.add_paragraph('Folder document conversion works.')
    doc.save(source / 'report.docx')
    workbook = Workbook()
    workbook.active.append(['Item', 'Quantity'])
    workbook.active.append(['Notebook', 3])
    workbook.save(nested / 'inventory.xlsx')
    workbook.close()
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((72,72), 'A folder can contain PDF documents.')
        pdf.save(source / 'book.pdf')
    (nested / 'notes.txt').write_text('Text files belong in the same queue.')
    (source / 'ignore.zip').write_bytes(b'unsupported')

    monkeypatch.setattr(app, 'default_preferences_path', lambda: tmp_path / 'preferences.json')
    monkeypatch.setattr(app, 'VAULT_DIR', None)
    api = app.Api()
    output = tmp_path / 'output'
    api.save_preferences({'output_dir': str(output)})
    api.stage_folder(source)
    assert len(api._collect_staged_paths()) == 4
    api._worker(api._collect_staged_paths())
    outputs = list(output.rglob('*.md'))
    assert len(outputs) == 4
    assert {path.parent.name for path in outputs} == {'docx', 'spreadsheets', 'pdf', 'txt'}
    content = '\n'.join(path.read_text() for path in outputs)
    for text in ('Folder document conversion works.', 'Notebook', 'A folder can contain PDF documents.', 'Text files belong in the same queue.'):
        assert text in content


def test_broken_nested_website_shortcut_does_not_block_valid_files(tmp_path, monkeypatch):
    app = load_converter_app()
    monkeypatch.setattr(app, 'default_preferences_path', lambda: tmp_path / 'settings' / 'preferences.json')
    source = tmp_path / 'source'
    nested = source / 'one' / 'two' / 'three'
    nested.mkdir(parents=True)
    valid = nested / 'notes.txt'
    valid.write_text('Keep this readable document even when its neighbor is broken.')
    broken = nested / 'Broken.webloc'
    broken.write_text('This is not a property list.')
    shortcut = nested / 'Useful.webloc'
    shortcut.write_bytes(plistlib.dumps({'URL': 'https://example.com/nested'}))
    api = app.Api()

    with mock.patch.object(api, '_log') as log:
        state = api.stage_folder(source)

    assert {item['source'] for item in state['items']} == {str(valid.resolve()), 'https://example.com/nested'}
    assert any(str(broken.resolve()) in call.args[0] and call.args[1] == 'log-error'
               for call in log.call_args_list)


@pytest.mark.skipif(os.geteuid() == 0, reason='Root can read directories even when permission bits are cleared.')
@pytest.mark.parametrize('readable_sibling', [False, True])
def test_unreadable_nested_directory_reports_path_without_false_empty_message(tmp_path, monkeypatch, readable_sibling):
    app = load_converter_app()
    monkeypatch.setattr(app, 'default_preferences_path', lambda: tmp_path / 'settings' / 'preferences.json')
    source = tmp_path / 'source'
    blocked = source / 'one' / 'two' / 'three'
    blocked.mkdir(parents=True)
    (blocked / 'notes.txt').write_text('The inaccessible directory must be reported.')
    original_mode = blocked.stat().st_mode & 0o777
    valid = source / 'available.txt'
    if readable_sibling:
        valid.write_text('A blocked subfolder must not prevent collecting this file.')
    api = app.Api()

    blocked.chmod(0)
    try:
        with mock.patch.object(api, '_log') as log:
            state = api.stage_files([str(source)])
    finally:
        blocked.chmod(original_mode)

    assert [item['source'] for item in state['items']] == ([str(valid.resolve())] if readable_sibling else [])
    assert any(str(blocked.resolve()) in call.args[0] and call.args[1] == 'log-error'
               for call in log.call_args_list)
    assert not any('No supported files found' in call.args[0] for call in log.call_args_list)


def test_deep_mixed_folder_restores_and_converts_distinct_duplicate_filenames(tmp_path, monkeypatch):
    import converter_app as app
    from docx import Document
    from openpyxl import Workbook
    from queue_runner import QueueRunner

    source = tmp_path / 'source'
    expected = {}
    branches = []
    for branch_name in ('alpha', 'beta'):
        branch = source / branch_name / 'two' / 'three'
        branch.mkdir(parents=True)
        branches.append(branch)
        note = branch / 'notes.TXT'
        note.write_text(f'{branch_name} nested text survives conversion.')
        expected[str(note.resolve())] = f'{branch_name} nested text'
        document = Document()
        document.add_paragraph(f'{branch_name} nested document survives conversion.')
        document_path = branch / 'report.docx'
        document.save(document_path)
        expected[str(document_path.resolve())] = f'{branch_name} nested document'
        workbook = Workbook()
        workbook.active.append(['Branch', 'Status'])
        workbook.active.append([f'{branch_name} nested spreadsheet', 'Ready'])
        workbook_path = branch / 'inventory.xlsx'
        workbook.save(workbook_path)
        workbook.close()
        expected[str(workbook_path.resolve())] = f'{branch_name} nested spreadsheet'
    (branches[0] / 'ignored.zip').write_bytes(b'unsupported')

    monkeypatch.setattr(app, 'default_preferences_path', lambda: tmp_path / 'settings' / 'preferences.json')
    monkeypatch.setattr(app, 'VAULT_DIR', None)
    api = app.Api()
    api.stage_folder(source)
    # Adding a nested folder and the same file separately must not duplicate work.
    api.stage_folder(branches[0])
    before_restart = api.stage_files([str(branches[1]), str(branches[0] / 'notes.TXT')])
    assert before_restart['total'] == len(expected)

    restored = app.Api()
    restored_state = restored.get_queue_state()
    assert restored_state['active_id'] == before_restart['active_id']
    assert {item['id'] for item in restored_state['items']} == {item['id'] for item in before_restart['items']}
    assert {item['source'] for item in restored_state['items']} == set(expected)
    assert restored_state['waiting'] == len(expected)
    output = tmp_path / 'output'
    runner = QueueRunner(restored._queue_store, tmp_path / 'settings' / 'queue-cache', {
        'output_dir': str(output), 'vault_dir': None, 'raw_ocr_mode': 'different',
    })
    runner.run(restored_state['active_id'])

    completed = restored.get_queue_state()
    assert completed['done'] == len(expected) and completed['failed'] == completed['waiting'] == 0
    assert len({item['output_path'] for item in completed['items']}) == len(expected)
    assert len(list(output.rglob('*.md'))) == len(expected)
    for item in completed['items']:
        destination = Path(item['output_path'])
        files = sorted(destination.glob('*.md')) if destination.is_dir() else [destination]
        assert files and all(path.is_file() for path in files)
        content = '\n'.join(path.read_text() for path in files)
        assert expected[item['source']] in content
