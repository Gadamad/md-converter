from pathlib import Path
import sys
import tempfile
import unittest
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
import converters
from preferences import Preferences
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


class ReviewRegressionTests(unittest.TestCase):
    def test_same_name_inputs_never_overwrite_output_or_vault(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            first=converters.write_output('first', 'report', 'a/report.txt', 1, root/'out', 'txt', root/'vault')
            second=converters.write_output('second', 'report', 'b/report.txt', 1, root/'out', 'txt', root/'vault')
            self.assertNotEqual(first.output_path, second.output_path)
            self.assertIn('first', Path(first.output_path).read_text())
            self.assertEqual(len(list((root/'vault'/'txt').glob('*.md'))), 2)

    def test_server_retry_delay_is_bounded_and_finite(self):
        for value in ['86400', 'inf', 'NaN']:
            response=mock.Mock(status_code=429, headers={'Retry-After': value})
            okay=mock.Mock(status_code=200, text='ok')
            with mock.patch.object(converters.requests, 'get', side_effect=[response, okay]), mock.patch.object(converters.time, 'sleep') as sleep:
                converters._fetch_url_html('https://example.com')
            delay=sleep.call_args.args[0]
            self.assertGreaterEqual(delay, 0)
            self.assertLessEqual(delay, converters.WEB_FETCH_MAX_DELAY_SECONDS)

    def test_invalid_preference_field_falls_back(self):
        for value in [[], {}, 42, False]:
            with self.subTest(value=value):
                self.assertIsNone(Preferences.from_dict({'output_dir': value}).output_dir)

    def test_docx_hyperlink_content_preserved_in_order(self):
        doc=Document()
        para=doc.add_paragraph('Before ')
        link=OxmlElement('w:hyperlink')
        run=OxmlElement('w:r')
        text=OxmlElement('w:t')
        text.text='linked important text'
        run.append(text)
        link.append(run)
        para._p.append(link)
        para.add_run(' after')
        self.assertEqual(converters._para_to_md(para), 'Before linked important text after')

    def test_docx_table_escapes_pipe_characters(self):
        doc=Document()
        table=doc.add_table(rows=1, cols=1)
        table.cell(0,0).text='a | b'
        self.assertIn('a \\| b', converters._table_to_md(table))


if __name__ == '__main__':
    unittest.main()
